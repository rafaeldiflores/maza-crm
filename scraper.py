"""
Motor principal de scraping de Google Maps para Talleres Multirubro.
Extrae nombres, teléfonos, enlaces directos a WhatsApp, direcciones, calificaciones y webs.
"""

import asyncio
import re
import sys
import urllib.parse
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

# Asegurar codificación UTF-8 en consolas Windows (cp1252)
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
try:  # Legado: el rastreo ahora es Places API (places.py). Opcional.
    from playwright.async_api import async_playwright, Page, BrowserContext
except ImportError:
    async_playwright = Page = BrowserContext = None

from config import BROWSER_CONFIG

def clean_phone(phone_str: Optional[str]) -> Optional[str]:
    """Limpia y estandariza el número de teléfono."""
    if not phone_str:
        return None
    # Remueve caracteres no numéricos excepto '+'
    cleaned = re.sub(r'[^\d+]', '', phone_str.strip())
    # Si viene en formato 9XXXXXXXX (chileno sin +56)
    if re.match(r'^9\d{8}$', cleaned):
        cleaned = '+56' + cleaned
    elif re.match(r'^569\d{8}$', cleaned):
        cleaned = '+' + cleaned
    return cleaned

def generate_whatsapp_link(phone_clean: Optional[str], taller_nombre: str = "") -> Optional[str]:
    """Genera un link directo a WhatsApp para contactar al taller."""
    if not phone_clean:
        return None
    # Eliminar el '+' para wa.me
    digits = re.sub(r'\D', '', phone_clean)
    if digits.startswith("569") and len(digits) == 11:
        mensaje = urllib.parse.quote(f"Hola, me gustaría consultar por los servicios de {taller_nombre}.")
        return f"https://wa.me/{digits}?text={mensaje}"
    return None

class WorkshopScraper:
    def __init__(self, headless: bool = True, max_results_per_query: int = 15, seen_keys: Optional[set] = None):
        self.headless = headless
        self.max_results = max_results_per_query
        self.results: List[Dict[str, Any]] = []
        # Si se entrega un set (ej. lo ya guardado en el CRM), lo usamos como
        # punto de partida: al re-escanear una comuna ya rastreada, los talleres
        # conocidos se saltan de inmediato (antes de hacer clic) en vez de
        # re-procesarlos completos cada vez.
        self._seen_keys = seen_keys if seen_keys is not None else set()
        self.skipped_known_count = 0

    async def _handle_consent(self, page: Page):
        """Maneja popups de consentimiento de Google si aparecen."""
        try:
            btn = page.locator("button:has-text('Rechazar todo'), button:has-text('Aceptar todo'), form[action*='consent'] button")
            if await btn.count() > 0:
                await btn.first.click(timeout=3000)
                await page.wait_for_timeout(1000)
        except Exception:
            pass

    async def scrape_query(self, context: BrowserContext, query: str, rubro_nombre: str, comuna: str, stop_event: Optional[asyncio.Event] = None) -> List[Dict[str, Any]]:
        """Ejecuta la búsqueda para un rubro y comuna específicos.

        `stop_event`, si se entrega, se revisa entre cada scroll y entre cada
        ficha procesada para poder cortar la extracción a mitad de camino
        (botón "Detener" de la UI) sin perder lo ya recolectado.
        """
        if stop_event and stop_event.is_set():
            return []
        page = await context.new_page()
        workshops = []
        
        full_query = f"{query} en {comuna} Chile"
        url = f"https://www.google.com/maps/search/{urllib.parse.quote(full_query)}"
        print(f"\n[+] Buscando: '{full_query}'...")
        
        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=25000)
            await self._handle_consent(page)
            
            # Verificar si existe panel de resultados
            feed = page.locator("div[role='feed']")
            try:
                await feed.wait_for(timeout=7000)
            except Exception:
                # Podría haber redirigido directo a un solo taller
                single_title = page.locator("h1.DUwDvf, h1")
                if await single_title.count() > 0:
                    t_name = await single_title.first.inner_text()
                    print(f"  -> Resultado único encontrado: {t_name}")
                    # Extraer directo
                    data = await self._extract_single_place(page, t_name, rubro_nombre, comuna)
                    if data:
                        workshops.append(data)
                else:
                    print("  [!] No se encontraron resultados o no cargó la lista.")
                await page.close()
                return workshops

            # Scroll adaptativo: en vez de 6 rondas fijas (a veces insuficientes,
            # a veces excesivas), scrollea hasta llegar al máximo pedido o hasta
            # que la lista deje de crecer durante 2 rondas seguidas.
            print(f"  -> Desplazando lista de resultados...")
            place_links = page.locator("div[role='feed'] a[href*='/maps/place/']")
            previous_count = -1
            stable_rounds = 0
            for _ in range(15):
                if stop_event and stop_event.is_set():
                    break
                current_count = await place_links.count()
                if current_count >= self.max_results:
                    break
                if current_count == previous_count:
                    stable_rounds += 1
                    if stable_rounds >= 2:
                        break
                else:
                    stable_rounds = 0
                previous_count = current_count
                await feed.evaluate("el => el.scrollBy(0, 1600)")
                await page.wait_for_timeout(650)

            place_links = page.locator("div[role='feed'] a[href*='/maps/place/']")
            total_places = await place_links.count()
            target_count = min(total_places, self.max_results)
            print(f"  -> Procesando {target_count} fichas de talleres encontradas...")

            for i in range(target_count):
                if stop_event and stop_event.is_set():
                    print("  [i] Extracción detenida por el usuario, cortando esta búsqueda.")
                    break

                link = place_links.nth(i)

                # Paso 1: datos mínimos desde la tarjeta de la lista. Si esto falla
                # no hay nada recuperable y se salta legítimamente el ítem.
                try:
                    href = await link.get_attribute("href") or page.url
                    card = link.locator("xpath=ancestor::div[contains(@class, 'Nv2PK') or @jsaction][1]")

                    nombre = ""
                    title_el = card.locator("div.qBF1Pd")
                    if await title_el.count() > 0:
                        nombre = await title_el.first.inner_text()
                    if not nombre:
                        nombre = await link.get_attribute("aria-label") or ""

                    nombre = nombre.strip()
                    if not nombre or nombre.lower() in ["resultados", "patrocinado", "anuncio"]:
                        continue

                    dedup_key = f"{nombre.lower()}_{comuna.lower()}"
                    if dedup_key in self._seen_keys:
                        self.skipped_known_count += 1
                        continue
                except Exception:
                    continue

                # Paso 2: abrir el panel de detalle y extraer campos ricos (teléfono,
                # dirección, rating, web). Si algo aquí falla (timeout de Maps,
                # selector desactualizado, etc.) igual conservamos el taller con
                # los datos base en vez de perderlo por completo.
                rating = ""
                opiniones = "0"
                phone_formatted = None
                wa_link = None
                direccion = ""
                website = ""

                try:
                    clicked = False
                    for attempt in range(2):
                        try:
                            await link.click(timeout=6000)
                            clicked = True
                            break
                        except Exception:
                            await page.wait_for_timeout(400)

                    if clicked:
                        # Esperar a que el panel de detalle realmente cargue en vez
                        # de un sleep fijo: más rápido cuando Maps responde rápido,
                        # más tolerante cuando responde lento.
                        try:
                            await page.locator("h1.DUwDvf, h1").first.wait_for(state="visible", timeout=6000)
                        except Exception:
                            pass

                        phone_btn = page.locator("button[data-item-id*='phone:tel:']")
                        addr_btn = page.locator("button[data-item-id*='address']")

                        # Dar un par de oportunidades extra si el panel sigue
                        # poblándose (teléfono/dirección suelen llegar después del título).
                        for _ in range(3):
                            if await phone_btn.count() > 0 or await addr_btn.count() > 0:
                                break
                            await page.wait_for_timeout(400)

                        rating_el = page.locator("span.ceNzKf, div.F7nice span[aria-hidden='true']")
                        rating = await rating_el.first.inner_text() if await rating_el.count() > 0 else ""

                        reviews_el = page.locator("span.RDApEe, div.F7nice span:nth-child(2)")
                        opiniones_raw = await reviews_el.first.inner_text() if await reviews_el.count() > 0 else ""
                        opiniones = re.sub(r'[^\d]', '', opiniones_raw) or "0"

                        phone = None
                        if await phone_btn.count() > 0:
                            raw_tel = await phone_btn.first.get_attribute("data-item-id")
                            phone = raw_tel.replace("phone:tel:", "").strip() if raw_tel else None
                        else:
                            try:
                                body_text = await page.locator("div[role='main']").inner_text()
                                phone_match = re.search(r'(\+?56\s?9\s?\d{4}\s?\d{4}|\b9\s?\d{4}\s?\d{4}\b|\b\d{2}\s?\d{4}\s?\d{4}\b)', body_text)
                                if phone_match:
                                    phone = phone_match.group(1)
                            except Exception:
                                pass

                        phone_formatted = clean_phone(phone)
                        wa_link = generate_whatsapp_link(phone_formatted, nombre)

                        direccion = await addr_btn.first.inner_text() if await addr_btn.count() > 0 else ""
                        direccion = direccion.strip()

                        web_btn = page.locator("a[data-item-id*='authority']")
                        website = await web_btn.first.get_attribute("href") if await web_btn.count() > 0 else ""
                except Exception as ex:
                    print(f"    [!] Detalle incompleto para '{nombre}': {ex}")

                entry = {
                    "Nombre": nombre,
                    "Rubro": rubro_nombre,
                    "Comuna": comuna,
                    "Teléfono": phone_formatted or "No especificado",
                    "WhatsApp Link": wa_link or "N/A",
                    "Dirección": direccion or "No especificada",
                    "Calificación": rating or "S/C",
                    "Opiniones": opiniones or "0",
                    "Sitio Web": website or "",
                    "Google Maps": href or "",
                    "Fecha Extracción": datetime.now().strftime("%Y-%m-%d %H:%M")
                }

                self._seen_keys.add(dedup_key)
                workshops.append(entry)
                print(f"    [OK] [{rubro_nombre}] {nombre} | Tel: {entry['Teléfono']}")

        except Exception as e:
            print(f"  [Error durante scraping]: {e}")
        finally:
            await page.close()

        return workshops

    async def scrape_batch(
        self,
        context: BrowserContext,
        jobs: List[Dict[str, str]],
        concurrency: int = 3,
        on_job_done=None,
        stop_event: Optional[asyncio.Event] = None
    ) -> List[Dict[str, Any]]:
        """Ejecuta varias búsquedas (query + rubro + comuna) en paralelo usando
        varias pestañas del mismo navegador, limitado por un semáforo.

        Antes cada combinación comuna×rubro×query se procesaba una por una de
        forma 100% secuencial, lo que hacía que un rastreo regional completo
        tomara horas. Con 3-4 pestañas concurrentes el tiempo total baja de
        forma casi lineal sin disparar bloqueos de Google Maps (similar a un
        usuario con varias pestañas abiertas a mano).

        Si `stop_event` se activa a mitad de camino (botón "Detener" de la UI),
        los trabajos que aún no empezaron se saltan de inmediato y los que ya
        están corriendo se cortan en el próximo punto de chequeo dentro de
        `scrape_query`, sin perder lo ya recolectado.
        """
        semaphore = asyncio.Semaphore(max(1, concurrency))
        all_results: List[Dict[str, Any]] = []

        async def run_job(job: Dict[str, str]):
            async with semaphore:
                if stop_event and stop_event.is_set():
                    res = []
                else:
                    try:
                        res = await self.scrape_query(
                            context=context,
                            query=job["query"],
                            rubro_nombre=job["rubro_nombre"],
                            comuna=job["comuna"],
                            stop_event=stop_event
                        )
                    except Exception as e:
                        print(f"  [Error en tarea de scraping] {job}: {e}")
                        res = []
                if on_job_done:
                    try:
                        await on_job_done(job, res)
                    except Exception as cb_err:
                        print(f"  [Error en callback de progreso]: {cb_err}")
                return res

        tasks = [asyncio.create_task(run_job(job)) for job in jobs]
        for finished in asyncio.as_completed(tasks):
            res = await finished
            all_results.extend(res)

        return all_results

    async def _extract_single_place(self, page: Page, nombre: str, rubro: str, comuna: str) -> Optional[Dict[str, Any]]:
        """Extrae detalles cuando Maps navega directamente a un único negocio."""
        try:
            phone_btn = page.locator("button[data-item-id*='phone:tel:']")
            phone = None
            if await phone_btn.count() > 0:
                raw_tel = await phone_btn.first.get_attribute("data-item-id")
                phone = raw_tel.replace("phone:tel:", "").strip()
            phone_formatted = clean_phone(phone)

            addr_btn = page.locator("button[data-item-id*='address']")
            direccion = await addr_btn.first.inner_text() if await addr_btn.count() > 0 else ""

            web_btn = page.locator("a[data-item-id*='authority']")
            website = await web_btn.first.get_attribute("href") if await web_btn.count() > 0 else ""

            return {
                "Nombre": nombre.strip(),
                "Rubro": rubro,
                "Comuna": comuna,
                "Teléfono": phone_formatted or "No especificado",
                "WhatsApp Link": generate_whatsapp_link(phone_formatted, nombre) or "N/A",
                "Dirección": direccion.strip() or "No especificada",
                "Calificación": "",
                "Opiniones": "0",
                "Sitio Web": website or "",
                "Google Maps": page.url,
                "Fecha Extracción": datetime.now().strftime("%Y-%m-%d %H:%M")
            }
        except Exception:
            return None

    def export_excel(self, data: List[Dict[str, Any]], filepath: Path):
        """Exporta los datos con formato ejecutivo para prospección."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Talleres Prospectos"

        headers = [
            "Nombre del Taller", "Rubro", "Comuna", "Teléfono", 
            "Link WhatsApp", "Dirección", "Estrellas", 
            "Reseñas", "Sitio Web", "Google Maps URL", "Fecha"
        ]
        ws.append(headers)

        # Estilos de encabezado
        header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="Segoe UI", size=10)
        link_font = Font(name="Segoe UI", size=10, color="0054E9", underline="single")
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB')
        )

        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 28

        for row_idx, item in enumerate(data, 2):
            ws.append([
                item["Nombre"],
                item["Rubro"],
                item["Comuna"],
                item["Teléfono"],
                item["WhatsApp Link"],
                item["Dirección"],
                item["Calificación"],
                item["Opiniones"],
                item["Sitio Web"],
                item["Google Maps"],
                item["Fecha Extracción"]
            ])

            ws.row_dimensions[row_idx].height = 20
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.font = cell_font
                cell.alignment = Alignment(vertical="center")

                # Formatear link de WhatsApp
                if col_idx == 5 and item["WhatsApp Link"] != "N/A":
                    cell.hyperlink = item["WhatsApp Link"]
                    cell.value = "Abrir WhatsApp"
                    cell.font = link_font
                elif col_idx == 10 and item["Google Maps"]:
                    cell.hyperlink = item["Google Maps"]
                    cell.value = "Ver en Maps"
                    cell.font = link_font

        # Ajuste automático del ancho de columnas
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(filepath)
        print(f"\n[OK] Archivo Excel guardado exitosamente en:\n    {filepath.resolve()}")

    def export_csv(self, data: List[Dict[str, Any]], filepath: Path):
        """Exporta en formato CSV con BOM UTF-8 para apertura perfecta en Excel en español."""
        import csv
        with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
            if not data:
                return
            writer = csv.DictWriter(f, fieldnames=list(data[0].keys()))
            writer.writeheader()
            writer.writerows(data)
        print(f"[OK] Archivo CSV guardado en: {filepath.resolve()}")
